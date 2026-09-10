"""Phase 2 exporter + summary + page-removal tests (mock provider)."""
from __future__ import annotations

from io import BytesIO

from tests.conftest import _auth, build_pptx, register_verified

PPTX_MIME = "application/vnd.openxmlformats-officedocument.presentationml.presentation"


def _paid_user_with_generated(client, n=2, amount=20):
    token, _ = register_verified(client)
    tp = client.post(
        "/api/billing/topup", headers=_auth(token),
        json={"amount": amount, "currency": "USD"},
    )
    assert tp.status_code == 200
    name, data = build_pptx(n)
    up = client.post(
        "/api/projects",
        headers=_auth(token),
        files={"file": (name, data, PPTX_MIME)},
    )
    assert up.status_code == 201, up.text
    pid = up.json()["id"]
    gen = client.post(f"/api/projects/{pid}/generate", headers=_auth(token))
    assert gen.status_code == 200, gen.text
    detail = client.get(f"/api/projects/{pid}", headers=_auth(token))
    assert detail.status_code == 200
    return token, pid, detail.json()


def test_pptx_writeback_contains_notes(client):
    token, pid, project = _paid_user_with_generated(client, 2)
    r = client.post(
        f"/api/projects/{pid}/export",
        headers=_auth(token),
        params={"fmt": "pptx", "strategy": "overwrite"},
    )
    assert r.status_code == 200, r.text
    from pptx import Presentation

    prs = Presentation(BytesIO(r.content))
    got = project["pages"]
    assert len(prs.slides) == 2
    for slide, page in zip(prs.slides, got):
        text = slide.notes_slide.notes_text_frame.text
        assert text == page["note_text"]


def test_docx_export(client):
    token, pid, _ = _paid_user_with_generated(client, 2)
    r = client.post(
        f"/api/projects/{pid}/export",
        headers=_auth(token),
        params={"fmt": "docx"},
    )
    assert r.status_code == 200
    assert "wordprocessingml" in r.headers["content-type"]
    assert "_notes.docx" in r.headers["content-disposition"]
    assert len(r.content) > 500  # a real zip/docx payload


def test_summary_and_delete_page(client):
    token, pid, _ = _paid_user_with_generated(client, 2)
    s = client.get(f"/api/projects/{pid}/summary", headers=_auth(token))
    assert s.status_code == 200, s.text
    body = s.json()
    assert body["total_chars"] > 0
    assert len(body["pages"]) == 2
    assert body["est_minutes"] > 0

    page_id = body["pages"][0]["ord"]
    page_db_id = _page_id_of_ord(client, token, pid, page_id)
    d = client.delete(
        f"/api/projects/{pid}/pages/{page_db_id}", headers=_auth(token)
    )
    assert d.status_code == 200
    assert d.json()["remaining"] == 1

    got = client.get(f"/api/projects/{pid}", headers=_auth(token)).json()
    assert len(got["pages"]) == 1


def _page_id_of_ord(client, token, pid, ord_no):
    p = client.get(f"/api/projects/{pid}", headers=_auth(token)).json()
    for page in p["pages"]:
        if page["ord"] == ord_no:
            return page["id"]
    raise AssertionError(f"no page ord {ord_no}")
