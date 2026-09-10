"""Generate a small 2-slide demo deck at the repo root for manual UI testing."""
from __future__ import annotations

from pathlib import Path

from pptx import Presentation

OUT = Path(__file__).resolve().parents[2] / "sample_deck.pptx"

SLIDES = [
    (
        "Why presentation prep is slow",
        [
            "Most speakers only draft notes the night before.",
            "Style and timing are inconsistent across decks.",
            "There is rarely time to practice the full talk.",
        ],
    ),
    (
        "How NotesBang helps",
        [
            "Upload the deck; notes are written per slide.",
            "Matched to your pace and target duration.",
            "Full script or cue cards, then export to PPTX.",
        ],
    ),
]


def main() -> int:
    prs = Presentation()
    for title, bullets in SLIDES:
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        if slide.shapes.title is not None:
            slide.shapes.title.text = title
        body = slide.placeholders[1]
        body.text = "\n".join(f"- {b}" for b in bullets)
    prs.save(str(OUT))
    print(f"wrote {OUT} ({OUT.stat().st_size} bytes, 2 slides)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
