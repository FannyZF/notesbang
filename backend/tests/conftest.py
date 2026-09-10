"""Shared pytest fixtures for Phase 0 backend tests."""
from __future__ import annotations

import os
import pathlib
import tempfile
import uuid
from io import BytesIO
from urllib.parse import parse_qs, urlparse

import pytest

# Configure an isolated SQLite DB BEFORE importing the app so the engine binds
# to the test database.
_test_db = pathlib.Path(tempfile.gettempdir()) / f"spekernotes_test_{os.getpid()}.db"
os.environ["DATABASE_URL"] = f"sqlite:///{_test_db}"
os.environ["ENVIRONMENT"] = "test"
os.environ["MAIL_DRIVER"] = "console"
os.environ["RATE_LIMIT_ENABLED"] = "false"
os.environ["PAGE_LIMIT"] = "80"
os.environ["TRIAL_PAGES_LIMIT"] = "2"

from fastapi.testclient import TestClient  # noqa: E402

from app.db.base import Base, engine  # noqa: E402
from app.main import app  # noqa: E402


@pytest.fixture(autouse=True)
def clean_db():
    Base.metadata.create_all(bind=engine)
    yield
    with engine.begin() as conn:
        for table in reversed(Base.metadata.sorted_tables):
            conn.execute(table.delete())


@pytest.fixture()
def client():
    with TestClient(app) as c:
        yield c


def register_verified(client: TestClient, email: str | None = None) -> tuple[str, str]:
    """Register + verify + login; returns (token, email)."""
    email = email or f"user-{uuid.uuid4().hex[:10]}@example.com"
    r = client.post(
        "/api/auth/register",
        json={"email": email, "password": "phase0secret"},
    )
    assert r.status_code == 201, r.text
    dev_url = r.json()["dev_verify_url"]
    assert dev_url is not None
    token = parse_qs(urlparse(dev_url).query)["token"][0]
    rv = client.get("/api/auth/verify", params={"token": token})
    assert rv.status_code == 200, rv.text
    rl = client.post(
        "/api/auth/login", json={"email": email, "password": "phase0secret"}
    )
    assert rl.status_code == 200, rl.text
    return rl.json()["token"], email


def _auth(token: str) -> dict:
    return {"Authorization": f"Bearer {token}"}


def build_pptx(n_slides: int = 2, prefix: str = "slide") -> tuple[str, bytes]:
    """Return (filename, bytes) of a tiny PPTX with ``n_slides`` text slides."""
    from pptx import Presentation

    prs = Presentation()
    for i in range(1, n_slides + 1):
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        title = slide.shapes.title
        if title is not None:
            title.text = f"{prefix} {i}"
        body = slide.placeholders[1]
        body.text = f"Content of {prefix} {i}.\nHello 世界，测试讲稿内容。"
    buf = BytesIO()
    prs.save(buf)
    buf.seek(0)
    return f"sample_{prefix}.pptx", buf.read()

