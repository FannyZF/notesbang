"""PPTX speaker-notes write-back (PRD §4.11 / §6 PageRevision semantics).

notes_strategy:
- overwrite: replace each slide's notes with the generated note.
- merge: keep the original notes and append the generated note below a marker.
"""
from __future__ import annotations

from io import BytesIO
from pathlib import Path

from pptx import Presentation

NOTES_MARKER = "（AI 生成讲稿）"


def write_pptx_notes(
    source_path: str | Path,
    pages: list[tuple[int, str]],
    strategy: str = "overwrite",
) -> bytes:
    prs = Presentation(str(source_path))
    for ord_num, note in pages:
        if ord_num < 1 or ord_num > len(prs.slides):
            continue
        slide = prs.slides[ord_num - 1]
        notes_frame = slide.notes_slide.notes_text_frame
        existing = notes_frame.text or ""
        if strategy == "merge" and existing.strip():
            combined = f"{existing}\n\n{NOTES_MARKER}\n{note}"
        else:
            combined = note
        notes_frame.text = combined

    buf = BytesIO()
    prs.save(buf)
    buf.seek(0)
    return buf.read()
