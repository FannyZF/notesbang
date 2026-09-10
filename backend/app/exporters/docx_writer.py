"""Word script export (PRD §4.11): per-slide title + note text."""
from __future__ import annotations

from io import BytesIO

from docx import Document


def write_docx(
    title: str,
    pages: list[tuple[str, str]],
) -> bytes:
    doc = Document()
    doc.add_heading(title or "Speaker Notes", level=0)
    for header, note in pages:
        doc.add_heading(header, level=2)
        if note:
            for para in note.splitlines():
                doc.add_paragraph(para)
    buf = BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.read()
