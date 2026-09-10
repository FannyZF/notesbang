"""Real DeepSeek end-to-end smoke over the HTTP API (PRD spikes).

Usage (from backend/, key via env only):
  $env:DEEPSEEK_API_KEY="sk-..." ; $env:LLM_PROVIDER="deepseek"
  python -m scripts.e2e_deepseek

Exercises: register+verify, upload pptx (trial), settings, plan, whole
generation, single-page regenerate, review, and export (pptx/docx/pdf).
"""
from __future__ import annotations

import os
import pathlib
import tempfile
from io import BytesIO
from urllib.parse import parse_qs, urlparse

_tmp = pathlib.Path(tempfile.gettempdir()) / "spekernotes_e2e.db"
if _tmp.exists():
    _tmp.unlink()
os.environ["DATABASE_URL"] = f"sqlite:///{_tmp}"
os.environ.setdefault("ENVIRONMENT", "e2e")
os.environ.setdefault("MAIL_DRIVER", "console")
os.environ.setdefault("EXEC_ASYNC", "false")

from fastapi.testclient import TestClient  # noqa: E402

from app.main import app  # noqa: E402

EMAIL = "e2e-deepseek@example.com"
PASSWORD = "e2edeepseek"

SAMPLE_DECK = [
    ("Introduction", "Today: turning a deck into a talk you can actually give."),
    ("The problem", "72% of presenters start the night before; quality suffers."),
    ("The idea", "Parse the deck, generate notes per slide, calibrate by pace."),
]


def build_pptx(n=3) -> bytes:
    from pptx import Presentation

    prs = Presentation()
    for i, (title, body) in enumerate(SAMPLE_DECK[:n], start=1):
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        if slide.shapes.title is not None:
            slide.shapes.title.text = title
        slide.placeholders[1].text = body
    buf = BytesIO()
    prs.save(buf)
    buf.seek(0)
    return buf.read()


def main() -> int:
    if not os.getenv("DEEPSEEK_API_KEY"):
        print("DEEPSEEK_API_KEY missing")
        return 2

    with TestClient(app) as client:
        r = client.post(
            "/api/auth/register",
            json={"email": EMAIL, "password": PASSWORD},
        )
        assert r.status_code == 201, r.text
        dev_url = r.json()["dev_verify_url"]
        token = parse_qs(urlparse(dev_url).query)["token"][0]
        assert client.get("/api/auth/verify", params={"token": token}).status_code == 200
        login = client.post(
            "/api/auth/login",
            json={"email": EMAIL, "password": PASSWORD},
        ).json()
        h = {"Authorization": f"Bearer {login['token']}"}

        data = build_pptx(2)
        up = client.post(
            "/api/projects",
            headers=h,
            files={"file": ("demo.pptx", data, "application/octet-stream")},
        )
        assert up.status_code == 201, up.text
        pid = up.json()["id"]

        setr = client.put(
            f"/api/projects/{pid}/settings",
            headers=h,
            json={
                "style": "business",
                "note_mode": "script",
                "target_minutes": 3,
                "output_lang": "en",
                "custom_scenario": "Pitch to a small investor group; persuasive and concrete.",
            },
        )
        assert setr.status_code == 200, setr.text

        plan = client.post(f"/api/projects/{pid}/plan", headers=h).json()
        print("PLAN total_units:", plan["total_units"],
              "| per page:", [p["target_chars"] for p in plan["pages"]])

        gen = client.post(f"/api/projects/{pid}/generate", headers=h)
        assert gen.status_code == 200, gen.text
        print("GENERATE:", gen.json())

        detail = client.get(f"/api/projects/{pid}", headers=h).json()
        for p in detail["pages"]:
            assert p["note_text"], f"empty notes on page {p['ord']}"
            print(f"  page {p['ord']}: {len(p['note_text'].split())} words -> "
                  f"{p['note_text'][:80]}...")

        # Single-page regenerate (trial budget allows one).
        regen = client.post(
            f"/api/projects/{pid}/pages/{detail['pages'][0]['id']}/regenerate",
            headers=h,
        )
        assert regen.status_code == 200, regen.text
        print("REGEN:", regen.json()["status"])

        review = client.post(f"/api/projects/{pid}/review", headers=h)
        assert review.status_code == 200, review.text
        print("REVIEW clean:", review.json()["clean"], "| issues:", len(review.json()["issues"]))

        # Trial lock on export must hold...
        locked = client.post(f"/api/projects/{pid}/export", headers=h)
        assert locked.status_code == 402, locked.text
        print("EXPORT trial lock: 402 OK")

        # ...then top up to unlock all formats.
        assert client.post(
            "/api/billing/topup", headers=h,
            json={"amount": 5, "currency": "USD"},
        ).status_code == 200
        for fmt in ("pptx", "docx", "pdf"):
            ex = client.post(
                f"/api/projects/{pid}/export", headers=h, params={"fmt": fmt}
            )
            assert ex.status_code == 200, ex.text
            out = pathlib.Path(tempfile.gettempdir()) / f"e2e_notes.{fmt}"
            out.write_bytes(ex.content)
            print(f"EXPORT {fmt}: {len(ex.content)} bytes -> {out.name}")

        report = client.get("/api/billing/report", headers=h).json()
        print("REPORT:", report)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
