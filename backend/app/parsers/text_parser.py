"""Plain-text document parsing for the content-scoring product.

Supports .docx (python-docx), .txt and .md. Returns (title, text).
"""
from __future__ import annotations

from io import BytesIO
from pathlib import Path


class ParseError(ValueError):
    pass


def _decode(data: bytes) -> str:
    for enc in ("utf-8", "utf-8-sig", "gb18030", "latin-1"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def parse_docx(data: bytes) -> str:
    from docx import Document as DocxDocument

    try:
        doc = DocxDocument(BytesIO(data))
    except Exception as exc:  # noqa: BLE001
        raise ParseError("Could not read the .docx file") from exc
    lines: list[str] = []
    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            lines.append(text)
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                lines.append(" | ".join(cells))
    return "\n".join(lines)


def parse_text(data: bytes) -> str:
    return _decode(data).strip()


def parse_upload(filename: str, data: bytes) -> tuple[str, str]:
    ext = Path(filename or "").suffix.lower()
    title = Path(filename or "Untitled").stem
    if ext == ".docx":
        text = parse_docx(data)
    elif ext in (".txt", ".md", ".markdown"):
        text = parse_text(data)
    elif ext == ".doc":
        raise ParseError("Legacy .doc is not supported; please save as .docx")
    else:
        raise ParseError("Supported formats: .docx, .txt, .md")
    if not text.strip():
        raise ParseError("The document appears to be empty")
    return title, text
